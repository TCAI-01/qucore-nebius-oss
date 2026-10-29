"""QuScan downloads: an Excel workbook and a PDF report built from a finished (or partly finished) scan."""
from __future__ import annotations

import io
import re
from html import escape

from .markdown_report import BRAND, BRAND_DARK, INK, LINE, MUTED, PANEL, _EMOJI
from .pdf_report import fonts

STATUS_LABEL = {"vulnerable": "Vulnerable", "at_risk": "At risk (unpinned)", "unknown": "No verdict",
                "safe": "Safe", "clear": "No known issues", "not_scanned": "Not scanned"}
SEV_HEX = {"critical": "E5484D", "high": "F07A3A", "medium": "E0A43A", "low": "3B9BF0", "none": "8B91A7"}
STATUS_HEX = {"vulnerable": "E5484D", "at_risk": "E0A43A", "unknown": "8B91A7", "safe": "3FB978", "clear": "2F9E6B", "not_scanned": "C3C9D4"}


def _cmds(row) -> str:
    return "\n".join(f"{c['label']}: {c['command']}" for c in row.get("commands") or [])


def _src(row) -> str:
    return "\n".join(f"[{i}] {s['url']}" for i, s in enumerate(row.get("evidence") or [], 1))


# ------------------------------------------------------------------- Excel ----

def write_xlsx(job: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    report = job.get("report") or {}
    book = Workbook()
    head_fill = PatternFill("solid", fgColor="095499")
    head_font = Font(bold=True, color="FFFFFF")
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Border(bottom=Side(style="thin", color="DCE3EC"))

    def sheet(title: str, headers: list[tuple[str, int]], rows: list[list], color_col: int | None = None, palette=None):
        ws = book.create_sheet(title)
        ws.append([h for h, _ in headers])
        for col, (_, width) in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col)
            cell.fill, cell.font, cell.alignment = head_fill, head_font, Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(col)].width = width
        for row in rows:
            ws.append(["" if v is None else v for v in row])
        for r in range(2, ws.max_row + 1):
            for c in range(1, len(headers) + 1):
                ws.cell(row=r, column=c).alignment = wrap
                ws.cell(row=r, column=c).border = thin
            if color_col is not None and palette:
                cell = ws.cell(row=r, column=color_col)
                key = str(cell.value or "").lower().replace(" (unpinned)", "").replace(" ", "_").replace("no_verdict", "unknown").replace("no_known_issues", "clear")
                color = palette.get(key)
                if color:
                    cell.fill = PatternFill("solid", fgColor=color)
                    cell.font = Font(bold=True, color="FFFFFF")
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
        return ws

    summary = book.active
    summary.title = "Summary"
    counts = job.get("counts") or {}
    lines = [("QuScan report", job.get("title", "")), ("Created", job.get("created_at", "")), ("Model", f"{job.get('deployment', '')} · {job.get('provider_label', '')}"),
             ("Risk score", f"{report.get('score', '')} / 100 (grade {report.get('grade') or '-'})"),
             ("Packages found", counts.get("packages", 0)), ("Packages scanned", counts.get("scanned", 0)),
             ("Vulnerable", counts.get("vulnerable", 0)), ("At risk (unpinned)", counts.get("at_risk", 0)),
             ("Safe", counts.get("safe", 0)), ("CVEs affecting the project", counts.get("cves", 0)),
             ("Critical / high", f"{counts.get('critical', 0)} / {counts.get('high', 0)}"),
             ("News alerts", counts.get("news", 0)), ("News alerts naming your packages", counts.get("affecting", 0)),
             ("Tavily requests", ", ".join(f"{k} {v}" for k, v in (report.get("tavily") or job.get("tavily") or {}).items())),
             ("Model tokens in / out", f"{(job.get('usage') or {}).get('input_tokens', 0)} / {(job.get('usage') or {}).get('output_tokens', 0)}")]
    for label, value in lines:
        summary.append([label, value])
    summary.column_dimensions["A"].width, summary.column_dimensions["B"].width = 34, 90
    summary["A1"].font = Font(bold=True, size=14, color="095499")
    for r in range(2, summary.max_row + 1):
        summary.cell(row=r, column=1).font = Font(bold=True)
    if (report.get("briefing") or {}).get("headline"):
        summary.append([])
        summary.append(["Today's headline", report["briefing"]["headline"]])
        summary.cell(row=summary.max_row, column=1).font = Font(bold=True)
        summary.cell(row=summary.max_row, column=2).alignment = wrap

    sheet("Alerts", [("Package", 28), ("Ecosystem", 11), ("Your version", 14), ("Status", 18), ("Severity", 11), ("Danger 0-10", 10),
                     ("CVEs", 30), ("Affected", 24), ("Upgrade to", 13), ("How to upgrade", 46), ("Summary", 70), ("Mitigation", 60), ("Sources", 60)],
          [[r["name"], r["ecosystem"], r["version"] or r["spec"] or "", STATUS_LABEL[r["status"]], r["severity"], r["danger_score"],
            ", ".join(c["id"] for c in r["cves"][:20]), "; ".join(r["affected_ranges"]), r["recommended_version"], _cmds(r),
            r["summary"], "\n".join(r["mitigation"]), _src(r)] for r in report.get("alerts") or []], 4, STATUS_HEX)
    sheet("CVEs", [("CVE", 18), ("Package", 26), ("Ecosystem", 11), ("Your version", 14), ("Affects you", 11), ("CVSS", 7), ("Severity", 10),
                   ("Affected", 24), ("Fixed in", 16), ("Updated", 12), ("Summary", 90), ("Link", 46)],
          [[c["id"], c["package"], c["ecosystem"], c.get("project_version") or "", "yes" if c["applies"] else "no", c.get("cvss"),
            c.get("severity"), c.get("affected"), c.get("fixed_in"), c.get("updated"), c.get("summary"), c.get("url")]
           for c in report.get("cves") or []], 7, SEV_HEX)
    sheet("Packages", [("Package", 30), ("Ecosystem", 11), ("Version", 14), ("Version from", 12), ("Declared range", 18), ("Direct", 8),
                       ("Status", 18), ("Severity", 10), ("CVEs", 7), ("Safe version", 13), ("Found in", 48), ("Verdict", 70)],
          [[r["name"], r["ecosystem"], r["version"] or "", r["version_from"], r["spec"] or "", "yes" if r["direct"] else "no",
            STATUS_LABEL[r["status"]], r["severity"], len(r["cves"]), r["safe_version"], "\n".join(r["sources"][:6]), r["reason"]]
           for r in report.get("packages") or []], 7, STATUS_HEX)
    sheet("Safe versions", [("Package", 30), ("Ecosystem", 11), ("Current", 16), ("Use version", 14), ("Status", 18), ("Basis", 90)],
          [[s["name"], s["ecosystem"], s["current"], s["safe_version"], STATUS_LABEL[s["status"]], s["basis"]] for s in report.get("safe") or []], 5, STATUS_HEX)
    briefing = report.get("briefing") or {}
    news_sources = briefing.get("sources") or []
    sheet("Threat news", [("Date", 12), ("Severity", 10), ("Category", 18), ("Title", 46), ("Affects your packages", 26), ("Packages named", 26),
                          ("CVEs", 24), ("Summary", 80), ("Action", 60), ("Sources", 60)],
          [[a["date"], a["severity"], a["category"], a["title"], ", ".join(a["affects_names"]), ", ".join(a["packages"]), ", ".join(a["cves"]),
            a["summary"], a["action"], "\n".join(f"[{n}] {news_sources[n - 1]['url']}" for n in a["sources"] if n <= len(news_sources))]
           for a in briefing.get("alerts") or []], 2, SEV_HEX)
    sheet("Manifests", [("File", 60), ("Kind", 20), ("Ecosystem", 11), ("Role", 12), ("Packages", 10)],
          [[m["path"], m["kind"], m["ecosystem"], m["role"], m["count"]] for m in report.get("manifests") or []])
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


# --------------------------------------------------------------------- PDF ----

def write_pdf(job: dict) -> bytes:
    from reportlab.graphics.shapes import Drawing, Rect, String
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import CondPageBreak, KeepTogether, LongTable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    regular, bold, unicode_font = fonts()
    report = job.get("report") or {}
    counts = job.get("counts") or {}
    width = A4[0] - 36 * mm
    base = ParagraphStyle("qBody", fontName=regular, fontSize=9, leading=13.5, textColor=INK, spaceAfter=5, splitLongWords=True)
    st = {
        "body": base, "small": ParagraphStyle("qSmall", parent=base, fontSize=7.6, leading=10.5, textColor=MUTED, spaceAfter=2),
        "title": ParagraphStyle("qTitle", parent=base, fontName=bold, fontSize=20, leading=25, textColor=BRAND, spaceAfter=4),
        "section": ParagraphStyle("qSection", parent=base, fontName=bold, fontSize=13.5, leading=18, textColor=BRAND, spaceBefore=10,
                                  spaceAfter=5, keepWithNext=True),
        "h3": ParagraphStyle("qH3", parent=base, fontName=bold, fontSize=10.2, leading=14, spaceBefore=6, spaceAfter=2, keepWithNext=True),
        "cell": ParagraphStyle("qCell", parent=base, fontSize=7.6, leading=10, spaceAfter=0),
        "head": ParagraphStyle("qHead", parent=base, fontName=bold, fontSize=7.6, leading=10, textColor=colors.white, spaceAfter=0),
        "tile": ParagraphStyle("qTile", parent=base, fontName=bold, fontSize=15, leading=19, spaceAfter=0),
        "label": ParagraphStyle("qLabel", parent=base, fontSize=7.2, leading=9.5, textColor=MUTED, spaceAfter=0),
        "code": ParagraphStyle("qCode", parent=base, fontName="Courier", fontSize=7.4, leading=10, spaceAfter=1),
    }

    def safe(text) -> str:
        text = _EMOJI.sub("", re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", " ", str(text if text is not None else "")))
        if not unicode_font:
            text = text.translate(str.maketrans({"—": "-", "–": "-", "“": '"', "”": '"', "’": "'", "‘": "'", "…": "...", "•": "-"}))
            text = text.encode("cp1252", "replace").decode("cp1252")
        return escape(text)

    def p(text, style="body"):
        return Paragraph(safe(text).replace("\n", "<br/>"), st[style])

    def table(rows, weights, colors_col=None, palette=None):
        total = sum(weights)
        data = [[Paragraph(safe(c), st["head"]) for c in rows[0]]] + [[c if isinstance(c, Paragraph) else Paragraph(safe(c), st["cell"]) for c in row] for row in rows[1:]]
        t = LongTable(data, colWidths=[width * w / total for w in weights], repeatRows=1)
        style = [("BACKGROUND", (0, 0), (-1, 0), BRAND), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                 ("LINEBELOW", (0, 1), (-1, -1), 0.4, LINE), ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                 ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
        if colors_col is not None and palette:
            for r, row in enumerate(rows[1:], 1):
                key = str(row[colors_col]).lower().replace(" (unpinned)", "").replace(" ", "_").replace("no_verdict", "unknown").replace("no_known_issues", "clear")
                if key in palette:
                    style.append(("BACKGROUND", (colors_col, r), (colors_col, r), colors.HexColor("#" + palette[key])))
                    style.append(("TEXTCOLOR", (colors_col, r), (colors_col, r), colors.white))
        t.setStyle(TableStyle(style))
        return t

    def tiles(items):
        cells = [[Paragraph(f'{safe(v)}', st["tile"]) for _, v in items], [Paragraph(safe(k), st["label"]) for k, _ in items]]
        t = Table(cells, colWidths=[width / len(items)] * len(items))
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), PANEL), ("BOX", (0, 0), (-1, -1), 0.5, LINE),
                               ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.white), ("LEFTPADDING", (0, 0), (-1, -1), 7),
                               ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
        return t

    def severity_bar(sev_counts):
        total = sum(sev_counts.values()) or 1
        drawing = Drawing(width, 26)
        x = 0.0
        for name in ("critical", "high", "medium", "low", "none"):
            value = sev_counts.get(name, 0)
            if not value:
                continue
            w = width * value / total
            drawing.add(Rect(x, 8, w, 14, fillColor=colors.HexColor("#" + SEV_HEX[name]), strokeColor=None))
            if w > 40:
                drawing.add(String(x + 4, 12, f"{name} {value}", fontName=bold, fontSize=7.5, fillColor=colors.white))
            x += w
        return drawing

    story = [p("QuScan · dependency threat report", "title"),
             p(f"{job.get('workspace_name') or ''} · {job.get('created_at', '')[:19].replace('T', ' ')} UTC · {job.get('deployment', '')} ({job.get('provider_label', '')}) · web data from Tavily", "small"),
             Spacer(1, 8)]
    story.append(tiles([("Risk score", f"{report.get('score') if report.get('score') is not None else '-'} {report.get('grade') or ''}"),
                        ("Scanned", f"{counts.get('scanned', 0)} / {counts.get('packages', 0)}"), ("Vulnerable", counts.get("vulnerable", 0)),
                        ("At risk", counts.get("at_risk", 0)), ("CVEs affecting you", counts.get("cves", 0)), ("News alerts", counts.get("news", 0))]))
    if any((report.get("severity_counts") or {}).values()):
        story += [Spacer(1, 6), p("Packages needing action, by severity", "small"), severity_bar(report["severity_counts"])]

    briefing = report.get("briefing") or {}
    if briefing.get("headline"):
        story += [p("Threat briefing", "section"), p(briefing["headline"], "h3")]

    alerts = report.get("alerts") or []
    story.append(p("Packages that need action", "section"))
    if not alerts:
        story.append(p("No scanned package is affected by a known vulnerability or compromise found in the sources."))
    for r in alerts:
        block = [p(f"{r['name']} {r['version'] or r['spec'] or ''}  ·  {STATUS_LABEL[r['status']]}  ·  {r['severity']}  ·  danger {r['danger_score']}/10", "h3"),
                 p(r["reason"], "small")]
        if r["summary"]:
            block.append(p(r["summary"]))
        if r["cves"]:
            block.append(p("CVEs: " + ", ".join(f"{c['id']} ({c.get('cvss') or 'n/a'})" for c in r["cves"][:12]), "small"))
        if r["recommended_version"]:
            block.append(p(f"Upgrade to {r['recommended_version']}", "body"))
            block += [Paragraph(safe(f"{c['label']}:  {c['command']}"), st["code"]) for c in r["commands"]]
        if r["mitigation"]:
            block += [p(f"- {m}", "small") for m in r["mitigation"]]
        if r["evidence"]:
            block.append(p("Sources: " + "  ".join(f"[{i}] {s['url']}" for i, s in enumerate(r["evidence"][:6], 1)), "small"))
        story += [CondPageBreak(60), KeepTogether(block[:3]), *block[3:], Spacer(1, 4)]

    cves = [c for c in report.get("cves") or [] if c["applies"]]
    if cves:
        story += [p("CVEs affecting the project", "section"),
                  table([["CVE", "Package", "Your version", "CVSS", "Severity", "Affected", "Fixed in", "Summary"]] +
                        [[c["id"], c["package"], c.get("project_version") or "", c.get("cvss") or "", c.get("severity") or "",
                          c.get("affected") or "", c.get("fixed_in") or "", (c.get("summary") or "")[:300]] for c in cves[:150]],
                        [13, 13, 9, 6, 8, 12, 9, 34], 4, SEV_HEX)]

    safe_rows = report.get("safe") or []
    if safe_rows:
        story += [p("Safe versions to use", "section"),
                  table([["Package", "Ecosystem", "Current", "Use version", "Status"]] +
                        [[s["name"], s["ecosystem"], s["current"], s["safe_version"] or "-", STATUS_LABEL[s["status"]]] for s in safe_rows],
                        [30, 12, 18, 18, 22], 4, STATUS_HEX)]

    if briefing.get("alerts"):
        sources = briefing.get("sources") or []
        story.append(p(f"Threat news · last {briefing.get('time_range', '')}", "section"))
        for a in briefing["alerts"]:
            refs = "  ".join(f"[{n}] {sources[n - 1]['url']}" for n in a["sources"] if n <= len(sources))
            story += [KeepTogether([p(f"{a['title']}  ·  {a['severity']}  ·  {a['category']}  ·  {a['date']}", "h3"),
                                    p(a["summary"]), p("Action: " + a["action"], "small") if a["action"] else Spacer(1, 1),
                                    p("Affects your packages: " + ", ".join(a["affects_names"]), "small") if a["affects_names"] else Spacer(1, 1),
                                    p(refs, "small") if refs else Spacer(1, 1)])]

    manifests = report.get("manifests") or []
    if manifests:
        story += [p("Dependency files read", "section"),
                  table([["File", "Kind", "Ecosystem", "Role", "Packages"]] + [[m["path"], m["kind"], m["ecosystem"], m["role"], m["count"]] for m in manifests],
                        [46, 18, 12, 12, 10])]
    story += [Spacer(1, 10), p("Findings come from Tavily web search, the OpenCVE database and the selected model's reading of those sources. "
                               "Version checks are computed from the affected ranges. Verify critical changes before deploying.", "small")]

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(regular, 7.5)
        canvas.setFillColor(MUTED)
        canvas.drawString(18 * mm, 10 * mm, "QuCoRe · QuScan")
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {doc.page}")
        canvas.setStrokeColor(BRAND_DARK)
        canvas.setLineWidth(2)
        canvas.line(18 * mm, A4[1] - 12 * mm, A4[0] - 18 * mm, A4[1] - 12 * mm)
        canvas.restoreState()

    stream = io.BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm,
                            title="QuScan report", author="QuCoRe")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()
